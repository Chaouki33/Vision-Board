"""Export des rapports en Word (.docx) et PDF, aux couleurs de Vision Board."""
import io
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

NAVY = (0x1B, 0x3A, 0x6B)
GOLD = (0xB8, 0x89, 0x2B)
INK = (0x14, 0x21, 0x3D)
MUTED = (0x5B, 0x6B, 0x82)
ASSESSMENT_COLORS = {"favorable": (0x1D, 0x7A, 0x3A), "à surveiller": (0x9A, 0x62, 0x00), "préoccupante": (0xB4, 0x23, 0x18)}
PRIORITY_LABELS = {"haute": "Priorité haute", "moyenne": "Priorité moyenne", "basse": "Priorité basse"}


def _period(meta: dict) -> str:
    start, end = meta.get("debut"), meta.get("fin")
    fr = lambda d: date.fromisoformat(d).strftime("%d/%m/%Y")  # noqa: E731
    if start and end:
        return f"Période du {fr(start)} au {fr(end)}"
    return "Ensemble des données disponibles"


def _author_note(report: dict) -> str:
    date_text = f"{report['created_at']:%d/%m/%Y}"
    if (report.get("model") or "").startswith("Vision Board"):
        return (f"Rapport généré automatiquement par Vision Board le {date_text} à partir des données importées "
                "(analyse par règles, sans IA). À relire avant diffusion.")
    return (f"Rapport rédigé par l'IA ({report['model']}) le {date_text} à partir des données importées dans "
            "Vision Board. À relire avant diffusion.")


def _subtitle(report: dict) -> str:
    ctx = report["context"]
    return f"{ctx['societe']['nom']} · {ctx['tableau_de_bord']['domaine']} · {_period(ctx['periode'])}"


# ---------------------------------------------------------------- Word
def _shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    props.append(shading)


def to_docx(report: dict) -> bytes:
    c = report["content"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        doc.styles[name].font.color.rgb = RGBColor(*NAVY)
        doc.styles[name].font.size = Pt(size)

    brand = doc.add_paragraph()
    run = brand.add_run("VISION BOARD")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10), RGBColor(*GOLD)
    title = doc.add_paragraph()
    run = title.add_run(c["titre"])
    run.bold, run.font.size, run.font.color.rgb = True, Pt(22), RGBColor(*NAVY)
    sub = doc.add_paragraph()
    run = sub.add_run(_subtitle(report))
    run.font.color.rgb = RGBColor(*MUTED)
    verdict = doc.add_paragraph()
    run = verdict.add_run(f"Appréciation globale : {c['appreciation_globale']}")
    run.bold, run.font.color.rgb = True, RGBColor(*ASSESSMENT_COLORS.get(c["appreciation_globale"], NAVY))

    doc.add_heading("Synthèse", level=1)
    doc.add_paragraph(c["synthese"])

    if c["chiffres_cles"]:
        doc.add_heading("Chiffres clés", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(table.rows[0].cells, ("Indicateur", "Valeur", "Commentaire")):
            cell.text = ""
            r = cell.paragraphs[0].add_run(text)
            r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, "1B3A6B")
        for item in c["chiffres_cles"]:
            cells = table.add_row().cells
            cells[0].text, cells[2].text = item["indicateur"], item["commentaire"]
            cells[1].text = ""
            cells[1].paragraphs[0].add_run(item["valeur"]).bold = True

    for section in c["sections"]:
        doc.add_heading(section["titre"], level=1)
        for paragraph in section["paragraphes"]:
            doc.add_paragraph(paragraph)

    for heading, items in (("Points forts", c["points_forts"]), ("Points de vigilance", c["points_de_vigilance"])):
        if items:
            doc.add_heading(heading, level=1)
            for item in items:
                doc.add_paragraph(item, style="List Bullet")

    if c["recommandations"]:
        doc.add_heading("Recommandations", level=1)
        for i, rec in enumerate(c["recommandations"], 1):
            p = doc.add_paragraph()
            head = p.add_run(f"{i}. {rec['action']}")
            head.bold = True
            meta = p.add_run(f"\n{PRIORITY_LABELS[rec['priorite']]} · horizon : {rec['horizon']}")
            meta.font.color.rgb, meta.font.size = RGBColor(*GOLD), Pt(9)
            doc.add_paragraph(rec["justification"])

    if c["limites_des_donnees"]:
        doc.add_heading("Limites de l'analyse", level=2)
        for item in c["limites_des_donnees"]:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(c["conclusion"])

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(_author_note(report))
    run.italic, run.font.size, run.font.color.rgb = True, Pt(8), RGBColor(*MUTED)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF
_FONT, _FONT_BOLD = "Helvetica", "Helvetica-Bold"
_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
]


def _register_fonts() -> None:
    """Police TrueType si disponible (accents et espaces insécables) ; Helvetica sinon."""
    global _FONT, _FONT_BOLD
    if _FONT != "Helvetica":
        return
    for regular, bold in _FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("VB", regular))
            pdfmetrics.registerFont(TTFont("VB-Bold", bold))
            # Associe <b> à la variante grasse dans les paragraphes.
            pdfmetrics.registerFontFamily("VB", normal="VB", bold="VB-Bold", italic="VB", boldItalic="VB-Bold")
            _FONT, _FONT_BOLD = "VB", "VB-Bold"
            return


def _text(value: str) -> str:
    text = value.replace("\u202f", " ").replace("\u00a0", " ")
    if _FONT == "Helvetica":
        text = text.replace("\u2212", "-").encode("cp1252", "replace").decode("cp1252")
    return escape(text)


def _rgb(c: tuple) -> colors.Color:
    return colors.Color(*(v / 255 for v in c))


def to_pdf(report: dict) -> bytes:
    _register_fonts()
    c = report["content"]
    styles = {
        "brand": ParagraphStyle("brand", fontName=_FONT_BOLD, fontSize=9, textColor=_rgb(GOLD), spaceAfter=4),
        "title": ParagraphStyle("title", fontName=_FONT_BOLD, fontSize=20, leading=24, textColor=_rgb(NAVY), spaceAfter=6),
        "sub": ParagraphStyle("sub", fontName=_FONT, fontSize=10, textColor=_rgb(MUTED), spaceAfter=6),
        "h1": ParagraphStyle("h1", fontName=_FONT_BOLD, fontSize=13.5, leading=17, textColor=_rgb(NAVY), spaceBefore=14, spaceAfter=6),
        "h2": ParagraphStyle("h2", fontName=_FONT_BOLD, fontSize=11.5, textColor=_rgb(NAVY), spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("body", fontName=_FONT, fontSize=10.5, leading=15, textColor=_rgb(INK), spaceAfter=6),
        "bullet": ParagraphStyle("bullet", fontName=_FONT, fontSize=10.5, leading=15, leftIndent=12, bulletIndent=2,
                                 textColor=_rgb(INK), spaceAfter=3),
        "meta": ParagraphStyle("meta", fontName=_FONT_BOLD, fontSize=8.5, textColor=_rgb(GOLD), spaceAfter=2),
        "cell": ParagraphStyle("cell", fontName=_FONT, fontSize=9.5, leading=12.5, textColor=_rgb(INK)),
        "cellh": ParagraphStyle("cellh", fontName=_FONT_BOLD, fontSize=9.5, textColor=colors.white),
        "note": ParagraphStyle("note", fontName=_FONT, fontSize=8, textColor=_rgb(MUTED), alignment=TA_CENTER, spaceBefore=18),
    }
    P = lambda text, style="body": Paragraph(_text(text), styles[style])  # noqa: E731

    story = [P("VISION BOARD", "brand"), P(c["titre"], "title"), P(_subtitle(report), "sub")]
    verdict_style = ParagraphStyle("verdict", parent=styles["body"], fontName=_FONT_BOLD,
                                   textColor=_rgb(ASSESSMENT_COLORS.get(c["appreciation_globale"], NAVY)))
    story.append(Paragraph(_text(f"Appréciation globale : {c['appreciation_globale']}"), verdict_style))
    story += [P("Synthèse", "h1"), P(c["synthese"])]

    if c["chiffres_cles"]:
        story.append(P("Chiffres clés", "h1"))
        rows = [[P("Indicateur", "cellh"), P("Valeur", "cellh"), P("Commentaire", "cellh")]]
        rows += [[P(i["indicateur"], "cell"), Paragraph(f"<b>{_text(i['valeur'])}</b>", styles["cell"]), P(i["commentaire"], "cell")]
                 for i in c["chiffres_cles"]]
        table = Table(rows, colWidths=[4.6 * cm, 3.6 * cm, 8.8 * cm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _rgb(NAVY)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.96, 0.98)]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.Color(0.86, 0.89, 0.93)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)

    for section in c["sections"]:
        story.append(P(section["titre"], "h1"))
        story += [P(p) for p in section["paragraphes"]]

    for heading, items in (("Points forts", c["points_forts"]), ("Points de vigilance", c["points_de_vigilance"])):
        if items:
            story.append(P(heading, "h1"))
            story += [Paragraph(_text(i), styles["bullet"], bulletText="•") for i in items]

    if c["recommandations"]:
        story.append(P("Recommandations", "h1"))
        for i, rec in enumerate(c["recommandations"], 1):
            story.append(KeepTogether([
                Paragraph(f"<b>{i}. {_text(rec['action'])}</b>", styles["body"]),
                P(f"{PRIORITY_LABELS[rec['priorite']]} · horizon : {rec['horizon']}", "meta"),
                P(rec["justification"]),
            ]))

    if c["limites_des_donnees"]:
        story.append(P("Limites de l'analyse", "h2"))
        story += [Paragraph(_text(i), styles["bullet"], bulletText="•") for i in c["limites_des_donnees"]]

    story += [P("Conclusion", "h1"), P(c["conclusion"])]
    story.append(P(_author_note(report), "note"))

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(_rgb(GOLD))
        canvas.setLineWidth(1.2)
        canvas.line(2 * cm, A4[1] - 1.3 * cm, A4[0] - 2 * cm, A4[1] - 1.3 * cm)
        canvas.setFont(_FONT, 8)
        canvas.setFillColor(_rgb(MUTED))
        canvas.drawString(2 * cm, 1.2 * cm, _text(report["context"]["societe"]["nom"]).replace("&amp;", "&"))
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Vision Board · page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=c["titre"], author="Vision Board").build(
        story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()
