"""Forme attendue des données de chaque domaine : description des colonnes, exemples et modèles de fichier.

Sert à guider l'utilisateur avant un import (quelles colonnes, quel format, à quoi elles servent),
à générer des modèles Excel / CSV prêts à remplir, et à la saisie directe dans l'application.
"""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalog import KPIS
from .fields import fields_for
from .references import DOMAINS

# Trois lignes d'exemple réalistes par domaine (valeurs Python : date, nombre ou texte).
D = dt.date
SAMPLE_ROWS: dict[str, list[dict]] = {
    "commercial": [
        {"date": D(2026, 1, 15), "client": "Marjane", "produit": "Huile 5L", "commercial": "Karim", "region": "Casablanca",
         "canal": "Grande distribution", "quantite": 130, "montant_ht": 12500, "cout_achat": 9200,
         "statut_paiement": "Réglé", "date_echeance": D(2026, 3, 15), "objectif": 250000},
        {"date": D(2026, 1, 16), "client": "Épicerie Nour", "produit": "Sucre 2kg", "commercial": "Salma",
         "region": "Agadir", "canal": "Détaillants", "quantite": 80, "montant_ht": 2240, "cout_achat": 1760,
         "statut_paiement": "Non réglé", "date_echeance": D(2026, 2, 15), "objectif": 250000},
        {"date": D(2026, 2, 3), "client": "Label'Vie", "produit": "Thé vert 500g", "commercial": "Karim",
         "region": "Rabat", "canal": "Grande distribution", "quantite": 60, "montant_ht": 2520, "cout_achat": 1620,
         "statut_paiement": "Réglé", "date_echeance": D(2026, 4, 3), "objectif": 260000},
    ],
    "production": [
        {"date": D(2026, 3, 2), "ligne": "Ligne A", "produit": "Biscuits 200g", "quantite_produite": 12400,
         "quantite_rebut": 310, "objectif": 13000, "temps_planifie": 16, "temps_arret": 1.5,
         "cadence_theorique": 950, "cout_production": 23500},
        {"date": D(2026, 3, 2), "ligne": "Ligne B", "produit": "Gaufrettes 100g", "quantite_produite": 16800,
         "quantite_rebut": 420, "objectif": 17000, "temps_planifie": 16, "temps_arret": 0.5,
         "cadence_theorique": 1200, "cout_production": 23500},
        {"date": D(2026, 3, 3), "ligne": "Ligne A", "produit": "Biscuits 200g", "quantite_produite": 11900,
         "quantite_rebut": 280, "objectif": 13000, "temps_planifie": 16, "temps_arret": 2,
         "cadence_theorique": 950, "cout_production": 22600},
    ],
    "stock": [
        {"date": D(2026, 9, 30), "article": "Farine T55 (sac 50 kg)", "categorie": "Matières premières",
         "entrepot": "Dépôt Casablanca", "quantite_stock": 120, "stock_minimum": 40, "cout_unitaire": 310,
         "entrees": 200, "sorties": 180, "date_peremption": D(2026, 12, 31)},
        {"date": D(2026, 9, 30), "article": "Cartons 40x30", "categorie": "Emballages", "entrepot": "Dépôt Casablanca",
         "quantite_stock": 350, "stock_minimum": 500, "cout_unitaire": 6, "entrees": 1000, "sorties": 1150,
         "date_peremption": None},
        {"date": D(2026, 9, 30), "article": "Biscuits 200g", "categorie": "Produits finis",
         "entrepot": "Entrepôt Bouskoura", "quantite_stock": 4200, "stock_minimum": 3000, "cout_unitaire": 4.2,
         "entrees": 52000, "sorties": 50500, "date_peremption": D(2027, 3, 31)},
    ],
    "rh": [
        {"matricule": "M001", "nom": "Amina Alaoui", "sexe": "F", "date_naissance": D(1990, 4, 12),
         "date_embauche": D(2019, 3, 1), "date_sortie": None, "poste": "Comptable", "departement": "Finance",
         "type_contrat": "CDI", "salaire_brut": 9500, "jours_absence": 2, "jours_conges_pris": 10,
         "jours_conges_acquis": 18, "heures_formation": 14},
        {"matricule": "M002", "nom": "Youssef Benali", "sexe": "H", "date_naissance": D(2002, 9, 3),
         "date_embauche": D(2026, 7, 1), "date_sortie": None, "poste": "Assistant commercial", "departement": "Commercial",
         "type_contrat": "Stage", "salaire_brut": 2500, "jours_absence": 0, "jours_conges_pris": 0,
         "jours_conges_acquis": 0, "heures_formation": 7},
        {"matricule": "M003", "nom": "Omar Kettani", "sexe": "H", "date_naissance": D(1985, 1, 20),
         "date_embauche": D(2015, 6, 15), "date_sortie": D(2026, 5, 31), "poste": "Magasinier",
         "departement": "Logistique", "type_contrat": "CDI", "salaire_brut": 4200, "jours_absence": 5,
         "jours_conges_pris": 18, "jours_conges_acquis": 18, "heures_formation": 0},
    ],
    "tresorerie": [
        {"date": D(2026, 9, 2), "banque": "Attijariwafa bank", "libelle": "Règlement client Marjane",
         "categorie": "Encaissements clients", "encaissement": 45000, "decaissement": None, "solde": 312000},
        {"date": D(2026, 9, 10), "banque": "Attijariwafa bank", "libelle": "Cotisations CNSS",
         "categorie": "Charges sociales", "encaissement": None, "decaissement": 18500, "solde": 293500},
        {"date": D(2026, 9, 28), "banque": "CIH Bank", "libelle": "Salaires de septembre",
         "categorie": "Salaires", "encaissement": None, "decaissement": 96000, "solde": 151000},
    ],
    "comptabilite": [
        {"compte": "7111", "libelle_compte": "Ventes de marchandises au Maroc", "debit": None, "credit": 2460000,
         "date": None, "journal": None},
        {"compte": "6111", "libelle_compte": "Achats de marchandises", "debit": 1795000, "credit": None,
         "date": None, "journal": None},
        {"compte": "5141", "libelle_compte": "Banques (soldes débiteurs)", "debit": 160000, "credit": None,
         "date": None, "journal": None},
    ],
}

FORMATS = {
    "date": "Date au format JJ/MM/AAAA",
    "number": "Nombre (ex. 12 500,50 — sans le symbole MAD)",
    "text": "Texte libre",
}

DESCRIPTIONS = {
    "commercial": "Une ligne par vente ou par facture (journal des ventes, export du logiciel de facturation).",
    "production": "Une ligne par jour et par ligne de production (ou par ordre de fabrication).",
    "stock": "Une ligne par article et par date d'inventaire (état de stock).",
    "rh": "Une ligne par salarié (liste du personnel, export de paie).",
    "tresorerie": "Une ligne par opération bancaire ou de caisse (relevé bancaire).",
    "comptabilite": "Une ligne par compte (balance générale) ou par écriture (grand livre).",
}


def _display(value, field_type: str) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        text = f"{value:,.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else f"{value:,}"
        return text.replace(",", " ").replace(".", ",")
    return str(value)


def field_guide(domain: str) -> dict:
    """Colonnes attendues : format, exemple, indicateurs alimentés et importance."""
    kpis = [k for k in KPIS if k.domain == domain]
    columns = []
    for field in fields_for(domain):
        users = [k for k in kpis if field.code in k.required]
        example = next((_display(r.get(field.code), field.type) for r in SAMPLE_ROWS[domain]
                        if r.get(field.code) is not None), "")
        columns.append({
            "code": field.code, "label": field.label, "type": field.type, "format": FORMATS[field.type],
            "example": example, "kpis": [k.name for k in users],
            # Essentielle : nécessaire à au moins un indicateur prioritaire du domaine.
            "essential": any(k.priority == 1 for k in users),
        })
    columns.sort(key=lambda c: (not c["essential"], -len(c["kpis"])))
    return {"domain": domain, "label": DOMAINS[domain], "description": DESCRIPTIONS[domain], "columns": columns,
            "examples": [{f.label: _display(r.get(f.code), f.type) for f in fields_for(domain)}
                         for r in SAMPLE_ROWS[domain]]}


def template_xlsx(domain: str) -> bytes:
    guide = field_guide(domain)
    fields = fields_for(domain)
    wb = Workbook()
    ws = wb.active
    ws.title = "Données"
    ws.append([f.label for f in fields])
    header_fill = PatternFill("solid", fgColor="1B3A6B")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row in SAMPLE_ROWS[domain]:
        ws.append([row.get(f.code) for f in fields])
    for i, f in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(f.label) + 4)
        for cell in ws[get_column_letter(i)][1:]:
            if f.type == "date":
                cell.number_format = "DD/MM/YYYY"
            elif f.type == "number":
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"

    help_ws = wb.create_sheet("Explications")
    help_ws.append([f"Modèle Vision Board — {guide['label']}"])
    help_ws["A1"].font = Font(bold=True, size=14, color="1B3A6B")
    help_ws.append([guide["description"] + " Remplacez les lignes d'exemple par vos données ; "
                    "les noms de colonnes peuvent être différents, ils sont reconnus automatiquement."])
    help_ws.append([])
    help_ws.append(["Colonne", "Importance", "Format attendu", "Exemple", "Indicateurs calculés grâce à elle"])
    for cell in help_ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for c in guide["columns"]:
        help_ws.append([c["label"], "Essentielle" if c["essential"] else "Facultative", c["format"], c["example"],
                        ", ".join(c["kpis"]) or "—"])
    for letter, width in zip("ABCDE", (32, 14, 42, 26, 90)):
        help_ws.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_csv(domain: str) -> bytes:
    fields = fields_for(domain)
    lines = [";".join(f.label for f in fields)]
    for row in SAMPLE_ROWS[domain]:
        lines.append(";".join(_display(row.get(f.code), f.type) for f in fields))
    return ("\n".join(lines) + "\n").encode("utf-8-sig")
