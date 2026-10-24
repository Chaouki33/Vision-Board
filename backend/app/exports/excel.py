"""Classeur Excel d'un tableau de bord, aux couleurs de l'application.

Feuilles : « Tableau de bord » (en-tête du secteur, tuiles des indicateurs, graphiques Excel natifs),
« Indicateurs » (détail et formules), « Graphiques » (séries des graphiques), puis une feuille
par fichier source avec ses lignes sur la période, sous forme de tableau Excel filtrable.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..reports.formatting import fmt_delta
from .data import PALETTE, STATUS_LABELS, ExportData

FONT = "Segoe UI"
STATUS_STYLE = {
    "ok": ("✓", PALETTE["ok"], PALETTE["ok_soft"]),
    "warning": ("!", PALETTE["warn"], PALETTE["warn_soft"]),
    "critical": ("✕", PALETTE["danger"], PALETTE["danger_soft"]),
}


def number_format(unit: str) -> str:
    if unit == "MAD":
        return '#,##0 "MAD"'
    if unit == "%":
        return '0.0 "%"'
    if unit == "ratio":
        return "0.00"
    if unit in ("jours", "mois", "heures", "ans"):
        return f'0.0 "{unit}"'
    return "#,##0.##"


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color.lstrip("#"))


def _font(color: str = PALETTE["text"], size: float = 10, bold: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color.lstrip("#"))


def _box(ws, top: int, left: int, bottom: int, right: int, fill: str, border: str | None = None) -> None:
    """Remplit un bloc de cellules et trace son contour."""
    side = Side(style="thin", color=border.lstrip("#")) if border else Side()
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            cell = ws.cell(r, c)
            cell.fill = _fill(fill)
            if border:
                cell.border = Border(left=side if c == left else None, right=side if c == right else None,
                                     top=side if r == top else None, bottom=side if r == bottom else None)


def _merged(ws, row: int, left: int, right: int, value, font: Font, align: Alignment | None = None):
    ws.merge_cells(start_row=row, start_column=left, end_row=row, end_column=right)
    cell = ws.cell(row, left, value)
    cell.font = font
    cell.alignment = align or Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    return cell


# ---------------------------------------------------------------- Feuille « Tableau de bord »
TILE_COLS = 4  # colonnes par tuile
TILES_PER_ROW = 3


def _tile_left(i: int) -> int:
    return 2 + i * (TILE_COLS + 1)  # B, G, L


def _dashboard_sheet(ws, data: ExportData, chart_refs: list[tuple]) -> None:
    s1, s2 = data.colors
    ws.title = "Tableau de bord"
    ws.sheet_view.showGridLines = False
    last_col = _tile_left(TILES_PER_ROW - 1) + TILE_COLS - 1
    ws.column_dimensions["A"].width = 2
    for i in range(TILES_PER_ROW):
        for c in range(TILE_COLS):
            ws.column_dimensions[get_column_letter(_tile_left(i) + c)].width = 10
        if i < TILES_PER_ROW - 1:
            ws.column_dimensions[get_column_letter(_tile_left(i) + TILE_COLS)].width = 2
    ws.column_dimensions[get_column_letter(last_col + 1)].width = 2

    # Fond de page sur la zone utile
    max_row = 8 + 5 * -(-len(data.kpis) // TILES_PER_ROW) + 16 * -(-len(data.charts) // 2) + 4
    for r in range(1, max_row + 1):
        for c in range(1, last_col + 2):
            ws.cell(r, c).fill = _fill(PALETTE["bg"])

    # En-tête aux couleurs du secteur
    ws.row_dimensions[1].height = 8
    _box(ws, 2, 2, 5, last_col, s1)
    for r in range(2, 6):
        ws.cell(r, last_col).fill = _fill(s2)  # liseré de la couleur secondaire
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 32
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 18
    _merged(ws, 2, 2, last_col - 1, f"VISION BOARD  ·  {data.company.upper()}", _font("#FFFFFF", 9, True))
    _merged(ws, 3, 2, last_col - 1, data.title, _font("#FFFFFF", 20, True))
    computed = sum(1 for k in data.kpis if k.error is None)
    _merged(ws, 4, 2, last_col - 1,
            f"{data.domain_label}  ·  {computed} indicateurs calculés sur {len(data.kpis)}  ·  {data.period_text}",
            _font("#FFFFFF", 10))
    tags = [data.sector_label, data.size_label, data.city, f"Exporté le {data.generated_at:%d/%m/%Y à %H:%M}"]
    _merged(ws, 5, 2, last_col - 1, "  ·  ".join(t for t in tags if t), _font("#FFFFFF", 9))
    ws.row_dimensions[6].height = 12

    # Tuiles des indicateurs
    row = 7
    for start in range(0, len(data.kpis), TILES_PER_ROW):
        heights = (30, 34, 20, 30)
        for offset, h in enumerate(heights):
            ws.row_dimensions[row + offset].height = h
        ws.row_dimensions[row + 4].height = 10
        for i, kpi in enumerate(data.kpis[start:start + TILES_PER_ROW]):
            left = _tile_left(i)
            right = left + TILE_COLS - 1
            status = STATUS_STYLE.get(kpi.status or "")
            _box(ws, row, left, row + 3, right, PALETTE["surface"], PALETTE["border"])
            if status:  # bande de couleur du statut, comme les cartes de l'application
                for r in range(row, row + 4):
                    ws.cell(r, left).border = Border(left=Side(style="thick", color=status[1].lstrip("#")),
                                                     top=ws.cell(r, left).border.top,
                                                     bottom=ws.cell(r, left).border.bottom)
            _merged(ws, row, left, right, kpi.name, _font(PALETTE["muted"], 9, True))
            value = _merged(ws, row + 1, left, right, kpi.value if kpi.value is not None else "—",
                            _font(PALETTE["primary"], 18, True))
            value.number_format = number_format(kpi.unit)
            if kpi.error:
                _merged(ws, row + 2, left, right, kpi.error, _font(PALETTE["muted"], 9))
            elif status:
                cell = _merged(ws, row + 2, left, left + 2, f"{status[0]} {STATUS_LABELS[kpi.status]}",
                               _font(status[1], 9, True))
                for c in range(left, left + 3):
                    ws.cell(row + 2, c).fill = _fill(status[2])
                cell.alignment = Alignment(vertical="center", indent=1)
            details = [t for t in (kpi.detail, fmt_delta(kpi.delta)) if t]
            _merged(ws, row + 3, left, right, "  ·  ".join(details), _font(PALETTE["muted"], 8))
        row += 5

    # Graphiques Excel natifs, deux par ligne
    row += 1
    for i, (kpi, cats, vals) in enumerate(chart_refs):
        anchor = f"{'B' if i % 2 == 0 else 'I'}{row + (i // 2) * 16}"
        ws.add_chart(_chart(kpi, cats, vals), anchor)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(last_col + 1)}{max_row}"


def _chart(kpi, categories: Reference, values: Reference):
    chart_data = kpi.chart
    horizontal = chart_data.type != "line"
    if horizontal:
        chart = BarChart()
        chart.type = "bar"
        chart.gapWidth = 60
        chart.x_axis.scaling.orientation = "maxMin"  # premier élément en haut, comme à l'écran
    else:
        chart = LineChart()
    chart.add_data(values, titles_from_data=False)
    chart.set_categories(categories)
    chart.title = f"{chart_data.title} ({chart_data.unit})" if chart_data.unit else chart_data.title
    chart.title.overlay = False
    chart.x_axis.tickLblPos = "low"  # libellés sous le graphique, même avec des valeurs négatives
    chart.legend = None
    chart.width, chart.height = 11.2, 7.4
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.numFmt = number_format(chart_data.unit)
    if horizontal:
        chart.y_axis.delete = True  # les valeurs sont écrites au bout des barres
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=PALETTE["border"].lstrip("#")))

    series = chart.series[0]
    color = PALETTE["series"].lstrip("#")
    if horizontal:
        series.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(noFill=True))
        for idx, label in enumerate(chart_data.labels):
            if label == "Autres":
                point = DataPoint(idx=idx)
                point.graphicalProperties = GraphicalProperties(solidFill=PALETTE["other"].lstrip("#"))
                series.dPt.append(point)
    else:
        series.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=color, w=28575))
        series.marker = Marker(symbol="circle", size=6)
        series.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill="FFFFFF"))
        series.smooth = False
    if horizontal:  # valeurs affichées au bout des barres, comme à l'écran
        series.dLbls = DataLabelList()
        series.dLbls.showVal = True
        for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(series.dLbls, attr, False)
        series.dLbls.numFmt = number_format(chart_data.unit)
    return chart


# ---------------------------------------------------------------- Autres feuilles
def _header(ws, row: int, labels: list[str]) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row, c, label)
        cell.font = _font("#FFFFFF", 10, True)
        cell.fill = _fill(PALETTE["primary"])
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _table(ws, name: str, rows: int, cols: int, style: str = "TableStyleMedium2") -> None:
    if rows == 0:
        return  # un tableau Excel doit contenir au moins une ligne
    table = Table(displayName=name, ref=f"A1:{get_column_letter(cols)}{rows + 1}")
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)


def _indicators_sheet(ws, data: ExportData) -> None:
    labels = ["Indicateur", "Valeur", "Unité", "Statut", "Variation vs mois précédent (%)", "Détail",
              "Formule", "Description", "Lignes utilisées", "Remarque"]
    _header(ws, 1, labels)
    for r, k in enumerate(data.kpis, start=2):
        values = [k.name, k.value, k.unit, STATUS_LABELS.get(k.status or "", ""), k.delta, k.detail,
                  k.formula, k.description, k.rows, k.error]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = _font()
            cell.alignment = Alignment(vertical="top", wrap_text=c in (6, 7, 8, 10))
        ws.cell(r, 2).number_format = number_format(k.unit)
        ws.cell(r, 5).number_format = '+0.0;-0.0;0'
        status = STATUS_STYLE.get(k.status or "")
        if status:
            ws.cell(r, 4).font = _font(status[1], 10, True)
    _table(ws, "Indicateurs", len(data.kpis), len(labels))
    for col, width in zip("ABCDEFGHIJ", (34, 16, 9, 14, 16, 36, 40, 50, 10, 24)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"


def _charts_sheet(ws, data: ExportData) -> list[tuple]:
    """Écrit les séries et renvoie, pour chaque graphique, les références des libellés et des valeurs."""
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    refs = []
    row = 1
    for kpi in data.charts:
        chart = kpi.chart
        ws.cell(row, 1, chart.title).font = _font(PALETTE["primary"], 11, True)
        _header(ws, row + 1, ["Libellé", f"Valeur ({chart.unit})" if chart.unit else "Valeur"])
        first = row + 2
        for i, (label, value) in enumerate(zip(chart.labels, chart.values)):
            ws.cell(first + i, 1, label).font = _font()
            cell = ws.cell(first + i, 2, value)
            cell.font = _font()
            cell.number_format = number_format(chart.unit)
        last = first + len(chart.values) - 1
        refs.append((kpi, Reference(ws, min_col=1, min_row=first, max_row=last),
                     Reference(ws, min_col=2, min_row=first, max_row=last)))
        row = last + 2
    return refs


def sheet_title(index: int, name: str) -> str:
    """Nom de feuille Excel valide (31 caractères au plus, sans caractère interdit), unique grâce au numéro."""
    base = re.sub(r"[\[\]:*?/\\']", " ", name).strip() or "Données"
    return f"{index} - {base}"[:31].strip()


def _unique_headers(columns: list[tuple[str, str]]) -> list[str]:
    seen: dict[str, int] = {}
    headers = []
    for name, _ in columns:
        label = str(name).strip() or "Colonne"
        if label.lower() in seen:
            seen[label.lower()] += 1
            label = f"{label} ({seen[label.lower()]})"
        seen.setdefault(label.lower(), 1)
        headers.append(label)
    return headers


def _source_sheet(ws, index: int, table) -> None:
    headers = _unique_headers(table.columns)
    _header(ws, 1, headers)
    formats = {"date": "dd/mm/yyyy", "number": "#,##0.##"}
    for r, values in enumerate(table.rows, start=2):
        for c, (value, (_, type_)) in enumerate(zip(values, table.columns), start=1):
            cell = ws.cell(r, c, value)
            if type_ in formats:
                cell.number_format = formats[type_]
    _table(ws, f"Source{index}", len(table.rows), len(headers), "TableStyleMedium9")
    for c, header in enumerate(headers, start=1):
        sample = [len(str(v[c - 1])) for v in table.rows[:200] if v[c - 1] is not None]
        ws.column_dimensions[get_column_letter(c)].width = min(max([len(header), *sample, 8]) + 3, 45)
    ws.freeze_panes = "A2"


def to_xlsx(data: ExportData) -> bytes:
    wb = Workbook()
    dashboard = wb.active
    indicators = wb.create_sheet("Indicateurs")
    charts = wb.create_sheet("Graphiques")
    _indicators_sheet(indicators, data)
    refs = _charts_sheet(charts, data)
    _dashboard_sheet(dashboard, data, refs)
    s1 = data.colors[0].lstrip("#")
    dashboard.sheet_properties.tabColor = s1
    indicators.sheet_properties.tabColor = PALETTE["primary"].lstrip("#")
    charts.sheet_properties.tabColor = PALETTE["primary"].lstrip("#")
    for i, table in enumerate(data.tables, start=1):
        ws = wb.create_sheet(sheet_title(i, table.name))
        ws.sheet_properties.tabColor = PALETTE["accent"].lstrip("#")
        _source_sheet(ws, i, table)
    wb.properties.title = f"{data.title} — {data.company}"
    wb.properties.creator = "Vision Board"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
