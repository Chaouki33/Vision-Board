import base64
import io
import json
import re
import zipfile
import zlib

from openpyxl import load_workbook

from .conftest import register

SALES = """Date facture;Client;Article;Commercial;Montant HT (MAD);Coût
15/01/2025;Marjane;Biscuits;Amine;12 500,00;8 000,00
20/01/2025;Label'Vie;Gaufrettes;Sara;8 400,00;5 100,00
10/02/2025;Marjane;Biscuits;Amine;15 000,00;9 200,00
03/03/2025;BIM;Cookies;Sara;9 800,00;6 300,00
""".encode("utf-8")


def _setup(client):
    headers, _ = register(client)
    company = client.post("/api/companies", headers=headers, json={
        "name": "Biscuiterie Export", "sector": "agroalimentaire", "employees": 40, "city": "Fès",
        "domains": ["commercial"]}).json()
    base = f"/api/companies/{company['id']}"
    r = client.post(f"{base}/datasets", headers=headers, data={"domain": "commercial"},
                    files={"file": ("ventes.csv", SALES, "text/csv")})
    assert r.status_code == 201, r.text
    dashboard = client.get(f"{base}/dashboards", headers=headers).json()[0]
    return headers, base, dashboard


def test_export_excel(client):
    headers, base, dashboard = _setup(client)
    r = client.get(f"{base}/dashboards/{dashboard['id']}/export?format=xlsx", headers=headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["Tableau de bord", "Indicateurs", "Graphiques"]
    assert wb["Tableau de bord"]._charts, "graphiques Excel attendus"
    assert wb["Tableau de bord"]["B3"].value == dashboard["title"]

    # Indicateurs : un par KPI du tableau de bord, avec la valeur calculée
    data = client.get(f"{base}/dashboards/{dashboard['id']}/data", headers=headers).json()
    indicators = wb["Indicateurs"]
    rows = list(indicators.iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == [k["name"] for k in data["kpis"]]
    assert [row[1] for row in rows] == [k["value"] for k in data["kpis"]]

    # Données sources : colonnes d'origine, montants en nombres, dates en dates
    source = wb[wb.sheetnames[3]]
    assert [c.value for c in source[1]] == ["Date facture", "Client", "Article", "Commercial", "Montant HT (MAD)", "Coût"]
    first = [c.value for c in source[2]]
    assert first[1] == "Marjane" and first[4] == 12500.0
    assert first[0].year == 2025 and first[0].month == 1
    assert source.max_row == 5


def test_export_period_filters_source_rows(client):
    headers, base, dashboard = _setup(client)
    r = client.get(f"{base}/dashboards/{dashboard['id']}/export?format=xlsx&start=2025-02-01&end=2025-03-31",
                   headers=headers)
    source = load_workbook(io.BytesIO(r.content)).worksheets[3]
    assert source.max_row == 3  # en-tête + février + mars


def _payload(expression: list[str]) -> list:
    packed = re.search(r'Binary\.FromText\("([^"]+)"', "\n".join(expression)).group(1)
    return json.loads(zlib.decompress(base64.b64decode(packed), -15))


def test_export_powerbi(client):
    headers, base, dashboard = _setup(client)
    r = client.get(f"{base}/dashboards/{dashboard['id']}/export?format=powerbi", headers=headers)
    assert r.status_code == 200, r.text
    archive = zipfile.ZipFile(io.BytesIO(r.content))
    names = archive.namelist()
    pbip = next(n for n in names if n.endswith(".pbip"))
    root, project = pbip.rsplit("/", 1)
    project = project[:-5]
    for part in (".Report/definition.pbir", ".Report/report.json", ".SemanticModel/definition.pbism",
                 ".SemanticModel/model.bim", ".xlsx"):
        assert f"{root}/{project}{part}" in names
    assert f"{root}/{project}.Report/StaticResources/RegisteredResources/VisionBoard.json" in names

    model = json.loads(archive.read(f"{root}/{project}.SemanticModel/model.bim"))["model"]
    tables = {t["name"]: t for t in model["tables"]}
    assert {"Indicateurs", "Graphiques"} <= tables.keys()

    # Les données intégrées se décompressent et reprennent les lignes du fichier
    source = next(t for name, t in tables.items() if name.startswith("Données"))
    rows = _payload(source["partitions"][0]["source"]["expression"])
    assert len(rows) == 4 and rows[0][2] == "Marjane" and rows[0][5] == "12500"
    assert rows[0][1] == "2025-01-15"
    data = client.get(f"{base}/dashboards/{dashboard['id']}/data", headers=headers).json()
    indicators = _payload(tables["Indicateurs"]["partitions"][0]["source"]["expression"])
    assert len(indicators) == len(data["kpis"])

    # Chaque champ utilisé par le rapport existe dans le modèle
    fields = {(t["name"], c["name"]) for t in model["tables"] for c in t["columns"]}
    fields |= {(t["name"], m["name"]) for t in model["tables"] for m in t.get("measures", [])}
    report = json.loads(archive.read(f"{root}/{project}.Report/report.json"))
    visuals = [json.loads(v["config"])["singleVisual"] for s in report["sections"] for v in s["visualContainers"]]
    used = 0
    for visual in visuals:
        query = visual.get("prototypeQuery")
        if not query:
            continue
        sources = {f["Name"]: f["Entity"] for f in query["From"]}
        for select in query["Select"]:
            item = select.get("Column") or select.get("Measure") or select["Aggregation"]["Expression"]["Column"]
            assert (sources[item["Expression"]["SourceRef"]["Source"]], item["Property"]) in fields
            used += 1
        refs = [p["queryRef"] for projection in visual["projections"].values() for p in projection]
        assert set(refs) <= {s["Name"] for s in query["Select"]}
    assert used > 0
    assert any(v["visualType"] == "card" for v in visuals)
    assert any(v["visualType"] in ("lineChart", "clusteredBarChart") for v in visuals)


def test_export_rejects_unknown_format_and_other_users(client):
    headers, base, dashboard = _setup(client)
    assert client.get(f"{base}/dashboards/{dashboard['id']}/export?format=pdf", headers=headers).status_code == 422
    other, _ = register(client)
    assert client.get(f"{base}/dashboards/{dashboard['id']}/export", headers=other).status_code == 404
