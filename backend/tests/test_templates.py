"""Forme attendue des données, modèles de fichier et saisie dans l'application."""
import io

import pytest
from openpyxl import load_workbook

from app.importers.mapping import suggest_mapping
from app.importers.tabular import parse_file
from app.kpi.fields import DOMAIN_FIELDS
from app.kpi.references import DOMAINS
from app.kpi.templates import field_guide, template_csv, template_xlsx


@pytest.mark.parametrize("domain", list(DOMAINS))
def test_templates_are_recognised_by_the_importer(domain):
    """Un modèle téléchargé puis réimporté tel quel est reconnu à 100 %."""
    for name, content in (("modele.xlsx", template_xlsx(domain)), ("modele.csv", template_csv(domain))):
        table = parse_file(name, content)
        mapping = suggest_mapping(domain, table.columns, table.rows)
        assert set(mapping) == {f.code for f in DOMAIN_FIELDS[domain]}, (domain, name, mapping)
        assert len(table.rows) == 3


def test_guide_describes_columns():
    guide = field_guide("commercial")
    by_code = {c["code"]: c for c in guide["columns"]}
    assert by_code["montant_ht"]["essential"] and "Chiffre d'affaires HT" in by_code["montant_ht"]["kpis"]
    assert by_code["date"]["format"].startswith("Date")
    assert by_code["montant_ht"]["example"] == "12 500"
    assert guide["columns"][0]["essential"]  # colonnes essentielles en premier
    wb = load_workbook(io.BytesIO(template_xlsx("rh")))
    assert wb.sheetnames == ["Données", "Explications"]


def test_template_endpoints(client, auth):
    assert client.get("/api/templates/stock", headers=auth).json()["label"] == "Gestion des stocks"
    r = client.get("/api/templates/stock/download", headers=auth, params={"format": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get("/api/templates/magie", headers=auth).status_code == 404


def test_manual_entry_create_edit_and_compute(client, auth, company):
    base = f"/api/companies/{company['id']}"
    rows = [
        {"date": "15/01/2026", "client": "Marjane", "montant_ht": "12 500,50", "statut_paiement": "Réglé"},
        {"date": "2026-02-10", "client": "BIM", "montant_ht": 4000, "statut_paiement": "Non réglé"},
        {"date": "", "client": "", "montant_ht": ""},  # ligne vide ignorée
    ]
    r = client.post(f"{base}/datasets/manual", headers=auth, json={"domain": "commercial", "name": "Ventes saisies", "rows": rows})
    assert r.status_code == 201, r.text
    dataset_id = r.json()["id"]
    assert r.json()["row_count"] == 2

    dash = client.post(f"{base}/dashboards", headers=auth,
                       json={"domain": "commercial", "title": "V", "kpi_codes": ["ca_total", "taux_impayes", "marge_brute"]}).json()
    kpis = {k["code"]: k for k in client.get(f"{base}/dashboards/{dash['id']}/data", headers=auth).json()["kpis"]}
    assert kpis["ca_total"]["value"] == 16500.5
    assert kpis["marge_brute"]["error"] == "Données manquantes"  # colonne laissée vide

    saved = client.get(f"{base}/datasets/{dataset_id}/rows", headers=auth).json()
    assert saved["editable"] and saved["rows"][0]["montant_ht"] == 12500.5 and saved["rows"][1]["date"] == "10/02/2026"
    saved["rows"][1]["montant_ht"] = 5000
    assert client.put(f"{base}/datasets/{dataset_id}/rows", headers=auth, json={"rows": saved["rows"]}).status_code == 200
    kpis = {k["code"]: k for k in client.get(f"{base}/dashboards/{dash['id']}/data", headers=auth).json()["kpis"]}
    assert kpis["ca_total"]["value"] == 17500.5


def test_manual_entry_reports_invalid_cells(client, auth, company):
    base = f"/api/companies/{company['id']}"
    r = client.post(f"{base}/datasets/manual", headers=auth, json={"domain": "commercial", "name": "X", "rows": [
        {"date": "31/02/2026", "montant_ht": "beaucoup"}]})
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert "Ligne 1, « Date de vente »" in detail and "« Montant HT » : nombre attendu" in detail
    empty = client.post(f"{base}/datasets/manual", headers=auth, json={"domain": "commercial", "name": "X", "rows": [{}]})
    assert empty.status_code == 422


def test_imported_files_are_not_editable(client, auth, company):
    base = f"/api/companies/{company['id']}"
    csv = "Date;Montant HT\n15/01/2026;100\n"
    ds = client.post(f"{base}/datasets", headers=auth, data={"domain": "commercial"},
                     files={"file": ("v.csv", csv.encode(), "text/csv")}).json()
    assert client.put(f"{base}/datasets/{ds['id']}/rows", headers=auth, json={"rows": [{"montant_ht": 1}]}).status_code == 409
