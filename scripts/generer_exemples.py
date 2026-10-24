"""Génère les fichiers de démonstration d'une société fictive par secteur d'activité.

Utilisation (depuis la racine du projet) :
    backend/.venv/Scripts/python scripts/generer_exemples.py

Les données sont fictives mais cohérentes entre elles : la balance comptable est construite à partir
des ventes, des salaires et des soldes bancaires générés pour la même société. Les tirages aléatoires
sont fixés : relancer le script redonne exactement les mêmes fichiers.
Atlas Distribution (commerce) et Atlas Biscuiterie (agroalimentaire) ont leurs propres fichiers.
"""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1] / "exemples"
START, END = dt.date(2026, 1, 2), dt.date(2026, 9, 30)
REF = dt.date(2026, 9, 30)
MONTHS = 9  # durée de la période des ventes et de la trésorerie

FIRST_F = ["Amina", "Salma", "Nadia", "Fatima Zahra", "Imane", "Khadija", "Sara", "Meryem", "Laila", "Zineb", "Hajar",
           "Asmae", "Kawtar", "Soukaina", "Hanane", "Ghizlane", "Najat", "Rim", "Houda", "Siham"]
FIRST_H = ["Youssef", "Karim", "Omar", "Hicham", "Rachid", "Mehdi", "Anas", "Hamza", "Ayoub", "Ilyas", "Adil", "Tarik",
           "Reda", "Soufiane", "Yassine", "Badr", "Abdellah", "Mustapha", "Nabil", "Othmane"]
LAST = ["Alaoui", "Benali", "Tazi", "Idrissi", "Fassi", "Berrada", "Chraibi", "Lahlou", "Amrani", "Squalli",
        "El Mansouri", "Bennani", "Kettani", "Naciri", "Ouazzani", "Sbai", "Cherkaoui", "Bouzidi", "El Amrani", "Haddad",
        "Ziani", "Rahmouni", "Benjelloun", "Saadi"]


def fr(x: float) -> str:
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def d(x: dt.date) -> str:
    return x.strftime("%d/%m/%Y")


def workdays(saturday: bool = True, sunday: bool = False):
    day = START
    while day <= END:
        if day.weekday() < 5 or (saturday and day.weekday() == 5) or (sunday and day.weekday() == 6):
            yield day
        day += dt.timedelta(days=1)


def write_csv(path: Path, header: str, lines: list[str]) -> None:
    path.write_text(header + "\n" + "\n".join(lines) + "\n", encoding="utf-8-sig")


# ---------------------------------------------------------------- Ventes
def ventes(folder: Path, s: dict, rng: random.Random) -> dict:
    """Retourne le CA et le coût des ventes de la période, pour construire la balance."""
    lines, n, ca, cost = [], 0, 0.0, 0.0
    days = list(workdays(saturday=s.get("saturday", True), sunday=s.get("sunday", False)))
    monthly_target = s["target"]
    for day in days:
        if rng.random() > s["sale_prob"]:
            continue
        for _ in range(rng.randint(*s["per_day"])):
            n += 1
            client = rng.choices(s["clients"], [c[1] for c in s["clients"]])[0][0]
            name, price, unit_cost, qmin, qmax = rng.choice(s["products"])
            q = rng.randint(qmin, qmax)
            amount, c = q * price * rng.uniform(0.95, 1.05), q * unit_cost
            ca, cost = ca + amount, cost + c
            unpaid = rng.random() < s["unpaid"]
            due = day + dt.timedelta(days=s["delay"])
            lines.append(f"{d(day)};{client};{name};{rng.choice(s['reps'])};{s['cities'].get(client, s['city'])};"
                         f"{s['channels'].get(client, 'Direct')};{q};{fr(amount)};{fr(c)};"
                         f"{'Non réglé' if unpaid else 'Réglé'};{d(due)};{monthly_target}")
    write_csv(folder / "ventes_2026.csv",
              "Date facture;Client;Produit;Commercial;Ville;Canal;Quantité;Montant HT;Coût de revient;"
              "Statut paiement;Échéance;Objectif mensuel", lines)
    return {"ca": ca, "cost": cost, "rows": n}


# ---------------------------------------------------------------- Production
def production(folder: Path, lines_spec: list, rng: random.Random) -> None:
    rows = []
    for day in workdays(saturday=True):
        for ligne, produit, cadence, planned, unit_cost, scrap, stop in lines_spec:
            arret = round(max(0.0, rng.gauss(stop, stop * 0.8)), 1) if rng.random() < 0.7 else 0.0
            arret = min(arret, planned - 1)
            output = int(cadence * (planned - arret) * rng.uniform(0.74, 0.94))
            rebut = int(output * rng.uniform(scrap * 0.5, scrap * 1.5))
            objectif = int(cadence * planned * 0.85)
            rows.append(f"{d(day)};{ligne};{produit};{output};{rebut};{objectif};{planned};{fr(arret)};{cadence};"
                        f"{fr(output * unit_cost * rng.uniform(0.96, 1.07))}")
    write_csv(folder / "production_2026.csv",
              "Date;Ligne;Produit;Quantité produite;Rebuts;Objectif;Temps planifié (h);Temps d'arrêt (h);"
              "Cadence théorique;Coût de production", rows)


# ---------------------------------------------------------------- Stock
def stock(folder: Path, articles: list, depots: list[str], rng: random.Random) -> float:
    """Deux inventaires (fin août, fin septembre). Retourne la valeur du dernier inventaire."""
    rows, value = [], 0.0
    for inv in (dt.date(2026, 8, 31), dt.date(2026, 9, 30)):
        value = 0.0
        for i, (name, family, pu, smin, perishable) in enumerate(articles):
            q = int(smin * rng.uniform(0.25, 3.6))
            if i == len(articles) - 1 and inv.month == 9:
                q = 0  # une rupture
            sorties = 0 if i == len(articles) - 2 else int(smin * rng.uniform(1.0, 5.0))  # un article dormant
            dlc = d(dt.date(2026, 10, 1) + dt.timedelta(days=rng.randint(-5, 240))) if perishable else ""
            rows.append(f"{d(inv)};{name};{family};{rng.choice(depots)};{q};{smin};{fr(pu)};{sorties};{dlc}")
            value += q * pu
    write_csv(folder / "stock_2026.csv",
              "Date inventaire;Article;Famille;Dépôt;Quantité en stock;Stock minimum;Prix unitaire;Sorties;"
              "Date de péremption", rows)
    return value


# ---------------------------------------------------------------- Personnel
def personnel(folder: Path, company: str, teams: list, women_share: float, rng: random.Random) -> float:
    """teams : (service, [postes], effectif, (salaire min, max), {contrat: poids}). Retourne la masse mensuelle."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel"
    ws.append([f"{company} — état du personnel au {d(REF)}"])
    ws.append([])
    ws.append(["Matricule", "Nom et prénom", "Sexe", "Date de naissance", "Date d'entrée", "Date de sortie",
               "Fonction", "Service", "Type de contrat", "Salaire brut", "Jours d'absence", "Congés acquis",
               "Congés pris", "Heures de formation"])
    n, payroll = 0, 0.0
    for service, posts, count, (lo, hi), contracts in teams:
        for _ in range(count):
            n += 1
            woman = rng.random() < women_share
            name = f"{rng.choice(FIRST_F if woman else FIRST_H)} {rng.choice(LAST)}"
            contract = rng.choices(list(contracts), list(contracts.values()))[0]
            hire = REF - dt.timedelta(days=rng.randint(20, 5000))
            if contract in ("Stage", "Saisonnier", "CDD"):
                hire = REF - dt.timedelta(days=rng.randint(15, 300))
            salary = 2500 if contract == "Stage" else round(rng.uniform(lo, hi), -1)
            leave = dt.date(2026, rng.randint(1, 8), rng.randint(1, 28)) if rng.random() < 0.08 else None
            if leave is None:
                payroll += salary
            ws.append([f"M{n:04d}", name, "F" if woman else "H",
                       dt.date(rng.randint(1966, 2004), rng.randint(1, 12), rng.randint(1, 28)), hire, leave,
                       rng.choice(posts), service, contract, salary, rng.choice([0, 0, 0, 1, 2, 3, 5, 8, 14]),
                       0 if contract in ("Stage", "Saisonnier") else 18,
                       0 if contract in ("Stage", "Saisonnier") else rng.randint(0, 18),
                       rng.choice([0, 0, 7, 14, 21, 35])])
    wb.save(folder / "personnel_2026.xlsx")
    return payroll


# ---------------------------------------------------------------- Trésorerie
def tresorerie(folder: Path, banks: dict, monthly_in: float, monthly_out: dict, rng: random.Random) -> float:
    """Encaissements clients et décaissements par catégorie, mois par mois. Retourne le solde final."""
    balances = dict(banks)
    rows = []
    names = list(banks)
    for month in range(START.month, END.month + 1):
        events = []
        for _ in range(rng.randint(10, 16)):
            events.append((rng.randint(2, 28), "in", "Règlement client", "Encaissements clients",
                           monthly_in / 13 * rng.uniform(0.6, 1.4)))
        for category, amount in monthly_out.items():
            parts = 1 if category in ("Salaires", "CNSS", "Loyer", "Impôts et taxes", "Échéance d'emprunt") else 4
            for _ in range(parts):
                day = 28 if category == "Salaires" else 10 if category == "CNSS" else rng.randint(3, 27)
                events.append((day, "out", category, category, amount / parts * rng.uniform(0.85, 1.15)))
        for day, kind, label, category, amount in sorted(events):
            bank = names[0] if kind == "out" and category in ("Salaires", "CNSS") else rng.choice(names)
            date = dt.date(2026, month, day)
            if kind == "in":
                balances[bank] += amount
                rows.append(f"{d(date)};{bank};{label};{category};{fr(amount)};;{fr(balances[bank])}")
            else:
                balances[bank] -= amount
                rows.append(f"{d(date)};{bank};{label};{category};;{fr(amount)};{fr(balances[bank])}")
    write_csv(folder / "tresorerie_2026.csv", "Date opération;Banque;Libellé;Catégorie;Crédit;Débit;Solde", rows)
    return sum(balances.values())


# ---------------------------------------------------------------- Balance comptable (exercice 2025)
def balance(folder: Path, company: str, b: dict) -> None:
    """Construit une balance équilibrée avant clôture à partir des grandeurs annuelles de la société."""
    D, C = "D", "C"
    ca = b["ca"]
    lines = []
    if b["merchandise"]:
        lines.append((7111, "Ventes de marchandises au Maroc", ca, C))
        lines.append((6111, "Achats de marchandises", ca * b["purchase_rate"], D))
    else:
        lines.append((7121 if b["goods"] else 7124, "Ventes de biens produits au Maroc" if b["goods"]
                      else "Prestations de services au Maroc", ca, C))
        if b["purchase_rate"]:
            lines.append((6121, "Achats de matières premières", ca * b["purchase_rate"], D))
    external = ca * b["external_rate"]
    for code, label, share in [(6125, "Achats non stockés de matières et fournitures (eau, électricité)", 0.25),
                               (6131, "Locations et charges locatives", 0.25), (6133, "Entretien et réparations", 0.15),
                               (6134, "Primes d'assurances", 0.08), (6142, "Transports", 0.12),
                               (6145, "Frais postaux et de télécommunications", 0.05),
                               (6144, "Publicité, publications et relations publiques", 0.10)]:
        lines.append((code, label, external * share, D))
    payroll = b["payroll"] * 12
    lines += [(6161, "Impôts et taxes d'exploitation", ca * 0.006, D),
              (6171, "Rémunérations du personnel", payroll, D),
              (6174, "Charges sociales", payroll * 0.21, D),
              (6193, "Dotations d'exploitation aux amortissements des immobilisations corporelles",
               b["fixed_assets"] * 0.08, D),
              (6311, "Intérêts des emprunts et dettes", b["loan"] * 0.065, D),
              (7381, "Intérêts et produits assimilés", b["cash"] * 0.004, C)]
    result = sum(a for _, _, a, s in lines if s == C) - sum(a for _, _, a, s in lines if s == D)
    lines.append((6701, "Impôts sur les bénéfices", max(result * 0.20, ca * 0.0025), D))

    fixed = b["fixed_assets"]
    lines += [(2321, "Bâtiments", fixed * b["buildings"], D),
              (2332, "Matériel et outillage", fixed * (0.85 - b["buildings"]), D),
              (2340, "Matériel de transport", fixed * 0.10, D), (2355, "Matériel informatique", fixed * 0.05, D),
              (2832, "Amortissements des constructions", fixed * b["buildings"] * 0.35, C),
              (2833, "Amortissements des installations techniques, matériel et outillage",
               fixed * (0.85 - b["buildings"]) * 0.45, C),
              (2834, "Amortissements du matériel de transport", fixed * 0.10 * 0.5, C),
              (2486, "Dépôts et cautionnements versés", ca * 0.004, D)]
    if b["stock"]:
        lines.append((3111 if b["merchandise"] else 3121, "Marchandises" if b["merchandise"] else "Matières premières",
                      b["stock"], D))
    purchases = ca * b["purchase_rate"] + external
    lines += [(3421, "Clients", ca * 1.2 * b["client_days"] / 360, D),
              (34552, "État – TVA récupérable sur charges", purchases * 0.2 * 0.08, D),
              (4411, "Fournisseurs", purchases * 1.2 * b["supplier_days"] / 360, C),
              (4432, "Rémunérations dues au personnel", b["payroll"] * 0.3, C),
              (4441, "CNSS", b["payroll"] * 0.21, C), (4455, "État – TVA facturée", ca * 0.2 * 0.08, C)]
    if b["cash"] >= 0:
        lines.append((5141, "Banques (soldes débiteurs)", b["cash"], D))
    else:
        lines.append((5541, "Banques (soldes créditeurs)", -b["cash"], C))
    lines += [(5161, "Caisses", ca * 0.0015, D), (1111, "Capital social", b["capital"], C),
              (1140, "Réserve légale", b["capital"] * 0.1, C)]
    if b["loan"]:
        lines.append((1481, "Emprunts auprès des établissements de crédit", b["loan"], C))
    lines = [(code, label, round(amount, 2), side) for code, label, amount, side in lines if round(amount, 2) > 0]
    gap = round(sum(a for *_, a, s in lines if s == D) - sum(a for *_, a, s in lines if s == C), 2)
    lines.append((1161, "Report à nouveau (solde créditeur)", gap, C) if gap > 0
                 else (1169, "Report à nouveau (solde débiteur)", -gap, D))
    lines.sort(key=lambda line: str(line[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Balance"
    ws.append([f"{company.upper()} — Balance générale au 31/12/2025 (avant inventaire)"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Compte", "Intitulé", "Débit", "Crédit"])
    for code, label, amount, side in lines:
        ws.append([code, label, amount if side == D else None, amount if side == C else None])
    total = round(sum(a for *_, a, s in lines if s == D), 2)
    ws.append(["Total général", None, total, total])
    ws.column_dimensions["B"].width = 62
    wb.save(folder / "balance_2025.xlsx")


# ---------------------------------------------------------------- Sociétés
def team(service, posts, count, salary, contracts=None):
    return (service, posts, count, salary, contracts or {"CDI": 80, "CDD": 15, "ANAPEC": 5})


SUPPORT = [team("Administration", ["Comptable", "Assistante de direction", "Responsable RH"], 4, (6000, 12000)),
           team("Direction", ["Directeur général", "Directeur administratif et financier"], 2, (18000, 35000),
                {"CDI": 100})]

COMPANIES = [
    {
        "slug": "maroc_metal_industries", "name": "Maroc Métal Industries SA", "sector": "industrie", "city": "Kénitra",
        "sales": {"clients": [("Stellantis Kénitra", 30), ("Renault Tanger", 25), ("Yazaki Morocco", 10),
                              ("Lear Corporation", 10), ("Bâtimat Distribution", 12), ("Export Espagne", 8)],
                  "products": [("Support moteur acier", 48, 34, 200, 1500), ("Châssis tubulaire", 420, 300, 10, 120),
                               ("Tôle découpée laser (lot)", 1250, 880, 5, 40), ("Pièce pliée inox", 95, 61, 100, 900),
                               ("Carter de protection", 160, 118, 50, 400)],
                  "reps": ["Mehdi", "Sanae", "Adil"], "city": "Kénitra",
                  "cities": {"Renault Tanger": "Tanger", "Bâtimat Distribution": "Casablanca", "Export Espagne": "Export"},
                  "channels": {"Stellantis Kénitra": "Automobile", "Renault Tanger": "Automobile", "Yazaki Morocco":
                               "Automobile", "Lear Corporation": "Automobile", "Bâtimat Distribution": "BTP",
                               "Export Espagne": "Export"},
                  "per_day": (1, 3), "sale_prob": 0.9, "unpaid": 0.12, "delay": 60, "target": 2_600_000},
        "production": [("Découpe laser", "Tôle découpée laser (lot)", 5, 16, 820, 0.03, 1.2),
                       ("Pliage", "Pièce pliée inox", 70, 16, 58, 0.04, 0.8),
                       ("Soudure", "Châssis tubulaire", 6, 16, 285, 0.06, 1.5),
                       ("Emboutissage", "Support moteur acier", 110, 16, 31, 0.025, 2.0)],
        "stock": ([("Tôle acier 2 mm (feuille)", "Matières premières", 640, 80, False),
                   ("Tube carré 40x40 (barre)", "Matières premières", 185, 150, False),
                   ("Inox 304 (feuille)", "Matières premières", 1450, 40, False),
                   ("Fil de soudure (bobine)", "Consommables", 320, 30, False),
                   ("Gaz argon (bouteille)", "Consommables", 540, 12, False),
                   ("Peinture époxy (seau)", "Consommables", 760, 20, True),
                   ("Boulonnerie (lot)", "Consommables", 95, 100, False),
                   ("Palettes bois", "Emballages", 70, 200, False),
                   ("Ancien moule presse", "Outillage", 4800, 1, False),
                   ("Disques de découpe", "Consommables", 38, 150, False)], ["Magasin central", "Atelier"]),
        "teams": [team("Production", ["Opérateur", "Soudeur", "Chef d'équipe", "Technicien maintenance"], 52,
                       (3800, 7500)), team("Qualité", ["Contrôleur qualité", "Responsable qualité"], 6, (5500, 13000)),
                  team("Logistique", ["Magasinier", "Cariste", "Chauffeur"], 8, (3700, 6000)),
                  team("Commercial", ["Ingénieur commercial", "ADV"], 5, (7000, 16000)),
                  team("Stagiaires", ["Stagiaire ingénieur"], 4, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.22, "banks": {"Attijariwafa bank": 1_200_000, "Société Générale Maroc": 450_000},
        "balance": {"merchandise": False, "goods": True, "purchase_rate": 0.52, "external_rate": 0.08,
                    "fixed_assets": 14_000_000, "buildings": 0.35, "loan": 4_500_000, "capital": 6_000_000,
                    "client_days": 70, "supplier_days": 75, "stock_months": 1.5},
    },
    {
        "slug": "tissage_du_nord", "name": "Tissage du Nord SARL", "sector": "textile", "city": "Tanger",
        "sales": {"clients": [("Inditex (sous-traitance)", 35), ("Mango", 15), ("Decathlon Maroc", 12),
                              ("Marjane (textile)", 10), ("Boutiques Defacto", 10), ("Export France", 8)],
                  "products": [("Jean denim", 115, 86, 150, 1200), ("Chemise coton", 68, 50, 200, 1500),
                               ("Veste de travail", 145, 104, 50, 400), ("T-shirt imprimé", 32, 23, 400, 2500)],
                  "reps": ["Hanane", "Tarik"], "city": "Tanger",
                  "cities": {"Decathlon Maroc": "Casablanca", "Marjane (textile)": "Rabat", "Export France": "Export"},
                  "channels": {"Inditex (sous-traitance)": "Sous-traitance export", "Mango": "Sous-traitance export",
                               "Export France": "Export"},
                  "per_day": (1, 3), "sale_prob": 0.85, "unpaid": 0.06, "delay": 45, "target": 1_900_000},
        "production": [("Coupe", "Jean denim", 140, 16, 78, 0.03, 0.6),
                       ("Couture A", "Chemise coton", 95, 16, 44, 0.05, 1.0),
                       ("Couture B", "Veste de travail", 40, 16, 93, 0.06, 1.2),
                       ("Sérigraphie", "T-shirt imprimé", 220, 16, 19, 0.08, 1.4)],
        "stock": ([("Tissu denim (mètre)", "Tissus", 38, 4000, False), ("Coton popeline (mètre)", "Tissus", 24, 5000, False),
                   ("Fil polyester (cône)", "Mercerie", 22, 600, False), ("Boutons métal (millier)", "Mercerie", 180, 40, False),
                   ("Fermetures éclair (centaine)", "Mercerie", 95, 80, False), ("Étiquettes tissées (millier)", "Mercerie", 140, 30, False),
                   ("Encre sérigraphie (kg)", "Consommables", 210, 50, True), ("Cartons d'expédition", "Emballages", 6, 3000, False),
                   ("Tissu imprimé collection 2024", "Tissus", 30, 100, False), ("Toile de travail (mètre)", "Tissus", 31, 2000, False)],
                  ["Entrepôt Tanger Free Zone"]),
        "teams": [team("Production", ["Couturière", "Coupeur", "Repasseuse", "Chef de chaîne", "Mécanicien machines"], 92,
                       (3300, 6200), {"CDI": 65, "CDD": 25, "ANAPEC": 10}),
                  team("Qualité", ["Contrôleuse qualité"], 8, (3800, 7000)),
                  team("Logistique", ["Magasinier", "Préparateur"], 7, (3500, 5500)),
                  team("Méthodes et modélisme", ["Modéliste", "Agent méthodes"], 5, (6000, 12000)),
                  team("Commercial", ["Chargée de clientèle export"], 3, (7000, 13000)),
                  team("Stagiaires", ["Stagiaire"], 3, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.72, "banks": {"BMCE Bank of Africa": 800_000, "CIH Bank": 300_000},
        "balance": {"merchandise": False, "goods": True, "purchase_rate": 0.56, "external_rate": 0.08,
                    "fixed_assets": 7_500_000, "buildings": 0.0, "loan": 2_000_000, "capital": 3_000_000,
                    "client_days": 55, "supplier_days": 60, "stock_months": 1.2},
    },
    {
        "slug": "batisseurs_du_souss", "name": "Bâtisseurs du Souss SARL", "sector": "btp", "city": "Agadir",
        "sales": {"clients": [("Al Omrane Souss-Massa", 25), ("Commune d'Agadir", 15), ("Promoteur Riad Souss", 25),
                              ("Ministère de l'Équipement", 10), ("Particuliers", 15), ("Hôtel Atlantic Bay", 10)],
                  "products": [("Situation de travaux gros œuvre", 185000, 152000, 1, 2),
                               ("Situation de travaux second œuvre", 96000, 76000, 1, 2),
                               ("Béton prêt à l'emploi (m³)", 820, 610, 20, 160),
                               ("Voirie et réseaux divers (lot)", 64000, 51000, 1, 2)],
                  "reps": ["Brahim", "Khalid"], "city": "Agadir",
                  "cities": {"Particuliers": "Inezgane", "Hôtel Atlantic Bay": "Taghazout"},
                  "channels": {"Al Omrane Souss-Massa": "Marchés publics", "Commune d'Agadir": "Marchés publics",
                               "Ministère de l'Équipement": "Marchés publics"},
                  "per_day": (1, 2), "sale_prob": 0.35, "unpaid": 0.28, "delay": 90, "target": 1_900_000},
        "production": [("Centrale à béton", "Béton prêt à l'emploi (m³)", 30, 10, 560, 0.02, 1.0),
                       ("Préfabrication", "Poutrelles précontraintes (ml)", 60, 10, 85, 0.04, 0.8),
                       ("Ferraillage", "Armatures façonnées (tonne)", 0.6, 10, 9800, 0.03, 0.5)],
        "stock": ([("Ciment CPJ 45 (sac 50 kg)", "Matériaux", 72, 800, False), ("Sable (m³)", "Matériaux", 180, 120, False),
                   ("Gravette (m³)", "Matériaux", 210, 120, False), ("Rond à béton Ø12 (tonne)", "Acier", 9200, 15, False),
                   ("Briques 8 trous (millier)", "Matériaux", 2100, 20, False), ("Adjuvant béton (bidon)", "Produits", 380, 25, True),
                   ("Coffrage métallique (m²)", "Matériel", 650, 50, False), ("Carrelage grès (m²)", "Second œuvre", 95, 600, False),
                   ("Ancienne bétonnière", "Matériel", 15000, 1, False), ("Tubes PVC Ø100 (barre)", "Réseaux", 64, 300, False)],
                  ["Dépôt Ait Melloul", "Chantier Founty"]),
        "teams": [team("Chantiers", ["Maçon", "Ferrailleur", "Chef de chantier", "Conducteur d'engins", "Manœuvre"], 40,
                       (3200, 9000), {"CDI": 45, "CDD": 50, "ANAPEC": 5}),
                  team("Bureau d'études", ["Ingénieur génie civil", "Métreur", "Dessinateur"], 6, (8000, 18000)),
                  team("Logistique", ["Magasinier", "Chauffeur poids lourd"], 6, (3800, 6500)),
                  team("Stagiaires", ["Stagiaire ingénieur"], 2, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.10, "banks": {"Banque Populaire": 350_000, "Crédit Agricole du Maroc": 150_000},
        "balance": {"merchandise": False, "goods": True, "purchase_rate": 0.55, "external_rate": 0.10,
                    "fixed_assets": 9_000_000, "buildings": 0.15, "loan": 3_500_000, "capital": 2_000_000,
                    "client_days": 125, "supplier_days": 95, "stock_months": 1.0},
    },
    {
        "slug": "conseil_audit_maghreb", "name": "Conseil & Audit Maghreb SARL", "sector": "services", "city": "Rabat",
        "sales": {"clients": [("Groupe OCP", 15), ("Banque régionale", 15), ("Ministère des Finances", 10),
                              ("Clinique Ibn Sina Privée", 10), ("PME industrielles (portefeuille)", 30), ("Association professionnelle", 10)],
                  "products": [("Mission d'audit (jour)", 7500, 3200, 3, 25), ("Conseil fiscal (jour)", 6500, 2800, 2, 10),
                               ("Formation (jour)", 9000, 3500, 1, 5), ("Accompagnement certification ISO (forfait)", 48000, 21000, 1, 1)],
                  "reps": ["Associé Benjelloun", "Associée Tazi"], "city": "Rabat",
                  "cities": {"Groupe OCP": "Casablanca", "Clinique Ibn Sina Privée": "Rabat"},
                  "channels": {"Ministère des Finances": "Marchés publics", "Association professionnelle": "Formation"},
                  "per_day": (1, 1), "sale_prob": 0.30, "unpaid": 0.18, "delay": 60, "target": 450_000},
        "teams": [team("Audit", ["Auditeur senior", "Auditeur junior", "Manager audit"], 8, (9000, 22000)),
                  team("Conseil", ["Consultant", "Consultant senior"], 5, (10000, 24000)),
                  team("Stagiaires", ["Stagiaire expertise comptable"], 3, (2500, 2500), {"Stage": 100}),
                  team("Administration", ["Office manager"], 1, (6000, 8000)),
                  team("Direction", ["Associé gérant"], 1, (35000, 40000), {"CDI": 100})],
        "women": 0.50, "banks": {"Attijariwafa bank": 420_000},
        "balance": {"merchandise": False, "goods": False, "purchase_rate": 0.0, "external_rate": 0.18,
                    "fixed_assets": 900_000, "buildings": 0.0, "loan": 0, "capital": 500_000,
                    "client_days": 85, "supplier_days": 45, "stock_months": 0},
    },
    {
        "slug": "digital_souk_technologies", "name": "Digital Souk Technologies SARL", "sector": "tic", "city": "Casablanca",
        "sales": {"clients": [("Clients SaaS PME (abonnements)", 40), ("Groupe de distribution", 15), ("Assurance Atlas", 15),
                              ("Startup fintech", 10), ("Administration publique", 10), ("Client export (France)", 10)],
                  "products": [("Abonnement SaaS (mensuel)", 1900, 380, 1, 12), ("Développement sur mesure (jour)", 5200, 2600, 5, 40),
                               ("Maintenance applicative (mois)", 14000, 6500, 1, 3), ("Hébergement cloud (mois)", 3500, 2100, 1, 6)],
                  "reps": ["Reda", "Kawtar"], "city": "Casablanca",
                  "cities": {"Client export (France)": "Export"},
                  "channels": {"Clients SaaS PME (abonnements)": "Abonnement", "Client export (France)": "Export"},
                  "per_day": (1, 2), "sale_prob": 0.55, "unpaid": 0.14, "delay": 45, "target": 1_050_000},
        "teams": [team("Développement", ["Développeur full-stack", "Développeur mobile", "Lead développeur"], 18,
                       (9000, 26000)), team("Infrastructure", ["Ingénieur DevOps", "Administrateur systèmes"], 4, (12000, 24000)),
                  team("Produit et design", ["Product owner", "Designer UX"], 4, (11000, 22000)),
                  team("Support client", ["Chargé de support"], 4, (5500, 8000)),
                  team("Commercial", ["Business developer"], 3, (9000, 18000)),
                  team("Stagiaires", ["Stagiaire développeur"], 4, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.35, "banks": {"CIH Bank": 900_000, "BMCE Bank of Africa": 250_000},
        "balance": {"merchandise": False, "goods": False, "purchase_rate": 0.0, "external_rate": 0.16,
                    "fixed_assets": 1_800_000, "buildings": 0.0, "loan": 600_000, "capital": 1_000_000,
                    "client_days": 65, "supplier_days": 40, "stock_months": 0},
    },
    {
        "slug": "trans_atlas_logistique", "name": "Trans Atlas Logistique SA", "sector": "transport", "city": "Tanger",
        "sales": {"clients": [("Renault Tanger", 20), ("Decathlon Maroc", 15), ("Lesieur Cristal", 15), ("OCP Group", 10),
                              ("Centrale Danone", 15), ("Transitaires portuaires", 15), ("Export Espagne", 10)],
                  "products": [("Transport complet Tanger–Casablanca", 6800, 5200, 1, 6), ("Transport frigorifique (trajet)", 8900, 6900, 1, 4),
                               ("Messagerie (palette)", 420, 300, 5, 60), ("Transport international Europe", 21000, 16500, 1, 2)],
                  "reps": ["Badr", "Siham"], "city": "Tanger",
                  "cities": {"Decathlon Maroc": "Casablanca", "Lesieur Cristal": "Casablanca", "OCP Group": "Jorf Lasfar",
                             "Centrale Danone": "Meknès", "Export Espagne": "Export"},
                  "channels": {"Export Espagne": "International", "Transitaires portuaires": "Port Tanger Med"},
                  "per_day": (2, 4), "sale_prob": 0.9, "unpaid": 0.16, "delay": 60, "target": 2_200_000},
        "teams": [team("Exploitation", ["Chauffeur poids lourd", "Chauffeur frigorifique", "Dispatcher"], 48, (4200, 8500)),
                  team("Atelier", ["Mécanicien poids lourds", "Chef d'atelier"], 7, (4500, 9500)),
                  team("Commercial", ["Commercial", "Responsable grands comptes"], 4, (8000, 17000)),
                  team("Stagiaires", ["Stagiaire logistique"], 2, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.08, "banks": {"Société Générale Maroc": 600_000, "Banque Populaire": 200_000},
        "balance": {"merchandise": False, "goods": False, "purchase_rate": 0.0, "external_rate": 0.42,
                    "fixed_assets": 22_000_000, "buildings": 0.05, "loan": 9_000_000, "capital": 5_000_000,
                    "client_days": 75, "supplier_days": 50, "stock_months": 0},
    },
    {
        "slug": "hotel_palmeraie_marrakech", "name": "Hôtel Palmeraie Marrakech SARL", "sector": "hotellerie_restauration",
        "city": "Marrakech",
        "sales": {"clients": [("Booking.com", 30), ("Expedia", 12), ("TUI (tour-opérateur)", 15), ("Clientèle directe", 20),
                              ("Entreprises (séminaires)", 13), ("Restaurant – clients extérieurs", 10)],
                  "products": [("Nuitée chambre double", 1450, 420, 2, 30), ("Nuitée suite", 3200, 780, 1, 6),
                               ("Restaurant (couvert)", 380, 145, 10, 120), ("Spa et hammam (prestation)", 650, 190, 2, 25),
                               ("Séminaire (journée)", 18000, 7600, 1, 1)],
                  "reps": ["Réception", "Service commercial"], "city": "Marrakech", "cities": {},
                  "channels": {"Booking.com": "Réservation en ligne", "Expedia": "Réservation en ligne",
                               "TUI (tour-opérateur)": "Tour-opérateur", "Entreprises (séminaires)": "Entreprises"},
                  "per_day": (2, 5), "sale_prob": 1.0, "sunday": True, "unpaid": 0.05, "delay": 30, "target": 1_700_000},
        "stock": ([("Viande bovine (kg)", "Restauration", 120, 60, True), ("Poulet fermier (kg)", "Restauration", 48, 80, True),
                   ("Huile d'olive (litre)", "Épicerie", 75, 40, True), ("Légumes frais (kg)", "Restauration", 9, 200, True),
                   ("Thé vert et menthe (kg)", "Épicerie", 120, 20, True), ("Eau minérale (pack)", "Boissons", 30, 150, True),
                   ("Jus d'orange frais (litre)", "Boissons", 14, 100, True), ("Linge de lit (parure)", "Hébergement", 450, 80, False),
                   ("Produits d'accueil (kit)", "Hébergement", 22, 600, False), ("Vaisselle de réception (ancien modèle)", "Restauration", 85, 10, False)],
                  ["Économat", "Cuisine centrale"]),
        "teams": [team("Hébergement", ["Réceptionniste", "Gouvernante", "Femme de chambre", "Bagagiste"], 18, (3500, 8000),
                       {"CDI": 70, "CDD": 25, "ANAPEC": 5}),
                  team("Restauration", ["Chef cuisinier", "Commis de cuisine", "Serveur", "Maître d'hôtel"], 17, (3500, 15000),
                       {"CDI": 65, "CDD": 30, "ANAPEC": 5}),
                  team("Spa", ["Esthéticienne", "Masseuse"], 5, (3800, 7000)),
                  team("Maintenance", ["Technicien maintenance", "Jardinier"], 5, (3600, 6500)),
                  team("Stagiaires", ["Stagiaire hôtellerie"], 4, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.48, "banks": {"Attijariwafa bank": 700_000, "CIH Bank": 180_000},
        "balance": {"merchandise": False, "goods": False, "purchase_rate": 0.0, "external_rate": 0.34,
                    "fixed_assets": 38_000_000, "buildings": 0.70, "loan": 15_000_000, "capital": 10_000_000,
                    "client_days": 20, "supplier_days": 50, "stock_months": 0.5},
    },
    {
        "slug": "pharmacie_al_amal", "name": "Pharmacie Centrale Al Amal", "sector": "sante", "city": "Fès",
        "sales": {"clients": [("Clients comptoir", 55), ("CNSS – AMO (tiers payant)", 15), ("CNOPS (tiers payant)", 12),
                              ("Mutuelles", 8), ("Cliniques du quartier", 10)],
                  "products": [("Médicaments (ordonnance)", 85, 66, 3, 40), ("Médicaments sans ordonnance", 42, 30, 2, 25),
                               ("Parapharmacie", 120, 78, 1, 10), ("Matériel médical", 260, 180, 1, 4), ("Lait infantile", 95, 76, 1, 12)],
                  "reps": ["Comptoir"], "city": "Fès", "cities": {},
                  "channels": {"CNSS – AMO (tiers payant)": "Tiers payant", "CNOPS (tiers payant)": "Tiers payant",
                               "Mutuelles": "Tiers payant", "Clients comptoir": "Comptoir"},
                  "per_day": (6, 12), "sale_prob": 1.0, "unpaid": 0.10, "delay": 75, "target": 360_000},
        "stock": ([("Amoxicilline 1 g (boîte)", "Médicaments", 38, 60, True), ("Paracétamol 500 mg (boîte)", "Médicaments", 9, 150, True),
                   ("Insuline (stylo)", "Médicaments", 210, 25, True), ("Sirop antitussif", "Médicaments", 32, 50, True),
                   ("Crème solaire SPF50", "Parapharmacie", 140, 30, True), ("Lait infantile 1er âge", "Nutrition", 92, 40, True),
                   ("Tensiomètre électronique", "Matériel médical", 380, 6, False), ("Masques chirurgicaux (boîte)", "Matériel médical", 35, 80, False),
                   ("Ancienne gamme cosmétique", "Parapharmacie", 95, 5, True), ("Vitamine D (ampoules)", "Médicaments", 24, 100, True)],
                  ["Officine"]),
        "teams": [team("Officine", ["Pharmacien adjoint", "Préparatrice", "Vendeur en parapharmacie"], 5, (4200, 12000)),
                  team("Stagiaires", ["Stagiaire en pharmacie"], 1, (2500, 2500), {"Stage": 100}),
                  team("Direction", ["Pharmacienne titulaire"], 1, (25000, 25000), {"CDI": 100})],
        "women": 0.70, "banks": {"Banque Populaire": 140_000},
        "balance": {"merchandise": True, "goods": False, "purchase_rate": 0.71, "external_rate": 0.05,
                    "fixed_assets": 650_000, "buildings": 0.0, "loan": 200_000, "capital": 300_000,
                    "client_days": 40, "supplier_days": 55, "stock_months": 1.8},
    },
    {
        "slug": "domaine_agrumes_berkane", "name": "Domaine Agrumes Berkane SARL", "sector": "agriculture", "city": "Berkane",
        "sales": {"clients": [("Exportateur Rungis (France)", 30), ("Marché de gros Casablanca", 20), ("Grande distribution Maroc", 15),
                              ("Export Russie", 15), ("Usine de jus (Agadir)", 10), ("Export Canada", 10)],
                  "products": [("Clémentines calibre 1 (kg)", 8.5, 4.6, 2000, 18000), ("Oranges Maroc Late (kg)", 5.2, 2.9, 3000, 25000),
                               ("Agrumes à jus (kg)", 2.1, 1.4, 5000, 30000)],
                  "reps": ["Service export", "Mustapha"], "city": "Berkane",
                  "cities": {"Marché de gros Casablanca": "Casablanca", "Usine de jus (Agadir)": "Agadir",
                             "Exportateur Rungis (France)": "Export", "Export Russie": "Export", "Export Canada": "Export"},
                  "channels": {"Exportateur Rungis (France)": "Export", "Export Russie": "Export", "Export Canada": "Export",
                               "Marché de gros Casablanca": "Marché de gros"},
                  "per_day": (1, 3), "sale_prob": 0.7, "unpaid": 0.09, "delay": 60, "target": 2_000_000},
        "production": [("Parcelle A – clémentiniers", "Clémentines calibre 1 (kg)", 900, 8, 4.1, 0.08, 0.5),
                       ("Parcelle B – orangers", "Oranges Maroc Late (kg)", 1200, 8, 2.6, 0.07, 0.5),
                       ("Station de conditionnement", "Agrumes conditionnés (kg)", 2500, 10, 0.9, 0.04, 1.2)],
        "stock": ([("Engrais NPK (sac 50 kg)", "Intrants", 420, 120, False), ("Produits phytosanitaires (bidon)", "Intrants", 680, 40, True),
                   ("Cartons export 10 kg", "Emballages", 7.5, 20000, False), ("Filets 2 kg", "Emballages", 1.2, 30000, False),
                   ("Cire d'enrobage (bidon)", "Conditionnement", 950, 20, True), ("Gasoil irrigation (litre)", "Énergie", 12.5, 5000, False),
                   ("Goutteurs irrigation (lot)", "Irrigation", 320, 50, False), ("Clémentines en chambre froide (kg)", "Produits", 4.6, 20000, True),
                   ("Ancien pulvérisateur", "Matériel", 9500, 1, False), ("Palettes export", "Emballages", 75, 800, False)],
                  ["Station Berkane", "Ferme Ain Reggada"]),
        "teams": [team("Exploitation agricole", ["Ouvrier agricole", "Cueilleur", "Chef d'équipe"], 170, (2900, 4200),
                       {"Saisonnier": 75, "CDI": 20, "CDD": 5}),
                  team("Station de conditionnement", ["Trieuse", "Opérateur calibreuse", "Cariste"], 45, (3000, 4800),
                       {"Saisonnier": 60, "CDI": 35, "CDD": 5}),
                  team("Technique", ["Ingénieur agronome", "Technicien irrigation"], 5, (7000, 15000)),
                  team("Export et commercial", ["Responsable export", "Agent commercial"], 3, (9000, 18000)),
                  team("Stagiaires", ["Stagiaire agronome"], 3, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.45, "banks": {"Crédit Agricole du Maroc": 1_500_000, "Attijariwafa bank": 500_000},
        "balance": {"merchandise": False, "goods": True, "purchase_rate": 0.30, "external_rate": 0.14,
                    "fixed_assets": 26_000_000, "buildings": 0.20, "loan": 8_000_000, "capital": 12_000_000,
                    "client_days": 55, "supplier_days": 70, "stock_months": 1.5},
    },
    {
        "slug": "immo_prestige_casablanca", "name": "Immo Prestige Casablanca SARL", "sector": "immobilier", "city": "Casablanca",
        "sales": {"clients": [("Acquéreurs particuliers", 45), ("Investisseurs MRE", 20), ("Entreprises (bureaux)", 15),
                              ("Propriétaires bailleurs (gestion)", 20)],
                  "products": [("Vente appartement (marge promotion)", 210000, 120000, 1, 1),
                               ("Commission de vente (agence)", 28000, 6000, 1, 1),
                               ("Gestion locative (mois)", 2400, 700, 1, 8), ("Vente plateau de bureaux (marge)", 380000, 230000, 1, 1)],
                  "reps": ["Nabil", "Houda", "Othmane"], "city": "Casablanca",
                  "cities": {"Investisseurs MRE": "Export"},
                  "channels": {"Propriétaires bailleurs (gestion)": "Gestion locative", "Investisseurs MRE": "MRE"},
                  "per_day": (1, 1), "sale_prob": 0.25, "unpaid": 0.08, "delay": 30, "target": 900_000},
        "teams": [team("Commercial", ["Conseiller immobilier", "Responsable d'agence"], 6, (6000, 16000)),
                  team("Gestion locative", ["Gestionnaire de biens"], 2, (6500, 9000)),
                  team("Stagiaires", ["Stagiaire commercial"], 1, (2500, 2500), {"Stage": 100})] + SUPPORT,
        "women": 0.45, "banks": {"BMCE Bank of Africa": 950_000, "CIH Bank": 300_000},
        "balance": {"merchandise": False, "goods": False, "purchase_rate": 0.0, "external_rate": 0.45,
                    "fixed_assets": 4_000_000, "buildings": 0.6, "loan": 1_500_000, "capital": 2_000_000,
                    "client_days": 30, "supplier_days": 45, "stock_months": 0},
    },
]


def generate(spec: dict) -> dict:
    rng = random.Random(spec["slug"])
    folder = ROOT / spec["slug"]
    folder.mkdir(parents=True, exist_ok=True)
    sales = ventes(folder, spec["sales"], rng)
    if "production" in spec:
        production(folder, spec["production"], rng)
    stock_value = stock(folder, *spec["stock"], rng) if "stock" in spec else 0
    payroll = personnel(folder, spec["name"], spec["teams"], spec["women"], rng)

    annual_ca = sales["ca"] * 12 / MONTHS
    b = spec["balance"]
    purchases = annual_ca * b["purchase_rate"]
    external = annual_ca * b["external_rate"]
    monthly_out = {"Salaires": payroll, "CNSS": payroll * 0.21,
                   "Fournisseurs": (purchases + external * 0.6) * 1.2 / 12,
                   "Loyer et charges": external * 0.25 / 12, "Énergie et télécoms": external * 0.15 / 12,
                   "Impôts et taxes": annual_ca * 0.01 / 12}
    if b["loan"]:
        monthly_out["Échéance d'emprunt"] = b["loan"] * 0.22 / 12
    # TVA reversée à l'État, impôt sur les sociétés (acomptes) et investissements courants.
    monthly_out["TVA à payer"] = max(0.0, annual_ca - purchases - external) * 0.20 / 12
    profit = (annual_ca - purchases - external - payroll * 12 * 1.21 - b["fixed_assets"] * 0.08
              - b["loan"] * 0.065)
    monthly_out["Impôt sur les sociétés"] = max(profit * 0.20, annual_ca * 0.0025) / 12
    monthly_out["Investissements"] = b["fixed_assets"] * 0.06 / 12
    cash = tresorerie(folder, spec["banks"], annual_ca * 1.2 / 12 * 0.97, monthly_out, rng)

    balance(folder, spec["name"], {**b, "ca": annual_ca, "payroll": payroll, "cash": cash,
                                   "stock": stock_value * 2 if stock_value else purchases / 12 * b["stock_months"]})
    return {"name": spec["name"], "ventes": sales["rows"], "ca_annuel": round(annual_ca), "masse_mensuelle": round(payroll),
            "tresorerie": round(cash)}


if __name__ == "__main__":
    for company in COMPANIES:
        print(generate(company))
